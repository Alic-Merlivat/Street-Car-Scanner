"""Renders /dashboard: a split-view HTML page reading straight from
detections.xlsx - a full data table on the left, and on the right, a
date-range picker plus the top 5 manufacturers and an hour-of-day timeline,
all scoped to whichever range is selected.

The left side has two tabs:
  - Cars: fully identified detections (known make/model and a readable plate).
  - Archive: everything else - unknown model, unreadable plate, or manually
    archived from the Cars tab - identified by the snapshot living under
    snapshots/review/ rather than snapshots/ directly.

No JS framework or charting library: bars are plain divs sized by percentage,
which keeps this dependency-free and easy to keep in sync with the sheet.
"""
import datetime as dt
import html
import os
from collections import Counter, defaultdict

from openpyxl import load_workbook

import config

_MULTI_WORD_MAKES = ("Mercedes", "Land Rover", "Alfa Romeo", "Aston Martin", "Rolls-Royce", "Range Rover")

_RANGE_LABELS = {
    "today": "Today",
    "yesterday": "Yesterday",
    "this_week": "This week",
    "last_week": "Last week",
    "this_month": "This month",
}


def _manufacturer(make_model: str | None) -> str | None:
    if not make_model or make_model == "unknown":
        return None
    for prefix in _MULTI_WORD_MAKES:
        if make_model.startswith(prefix):
            return "Mercedes-Benz" if prefix == "Mercedes" else prefix
    return make_model.split(" ", 1)[0]


def _date_bounds(preset: str) -> tuple[dt.date, dt.date]:
    today = dt.date.today()
    if preset == "yesterday":
        d = today - dt.timedelta(days=1)
        return d, d
    if preset == "this_week":
        start = today - dt.timedelta(days=today.weekday())  # Monday
        return start, today
    if preset == "last_week":
        this_monday = today - dt.timedelta(days=today.weekday())
        start = this_monday - dt.timedelta(days=7)
        return start, this_monday - dt.timedelta(days=1)
    if preset == "this_month":
        return today.replace(day=1), today
    return today, today  # "today" (and fallback default)


def _read_rows() -> list[dict]:
    if not config.EXCEL_PATH.exists():
        return []
    workbook = load_workbook(config.EXCEL_PATH, read_only=True)
    sheet = workbook["Detections"]
    rows = []
    for r in sheet.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        rows.append({
            "date": str(r[0]),
            "time": str(r[1]),
            "vehicle_type": r[2],
            "make_model": r[3],
            "plate": r[5],
            "snapshot": r[7],
        })
    return rows


def _in_range(r: dict, start: dt.date, end: dt.date) -> bool:
    try:
        d = dt.date.fromisoformat(r["date"])
    except ValueError:
        return False
    return start <= d <= end


def _is_identified(r: dict) -> bool:
    return bool(
        r["make_model"] and r["make_model"] != "unknown"
        and r["plate"] and r["plate"] != "unreadable"
    )


def _is_archived(r: dict) -> bool:
    snapshot = r["snapshot"] or ""
    return not _is_identified(r) or snapshot.startswith("review/")


def archive_row(snapshot_filename: str) -> tuple[bool, str | None]:
    """Moves a Cars-tab detection into the archive: the image file goes into
    snapshots/review, and the sheet's Snapshot File cell + hyperlink are
    updated to match, which is what makes it show up under the Archive tab.

    Returns (success, error_code). error_code is "locked" if detections.xlsx
    is open in another program (e.g. Excel) - the file move is rolled back
    in that case rather than left half-done."""
    src = config.SNAPSHOT_DIR / snapshot_filename
    if not src.exists() or not config.EXCEL_PATH.exists():
        return False, "not_found"

    try:
        workbook = load_workbook(config.EXCEL_PATH)
    except PermissionError:
        return False, "locked"

    sheet = workbook["Detections"]
    target_cell = None
    for row in sheet.iter_rows(min_row=2):
        cell = row[7]  # Snapshot File column
        if cell.value == snapshot_filename:
            target_cell = cell
            break
    if target_cell is None:
        return False, "not_found"

    config.REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    dest = config.REVIEW_DIR / snapshot_filename
    os.replace(src, dest)
    target_cell.value = f"review/{snapshot_filename}"
    target_cell.hyperlink = dest.as_uri()

    tmp_path = config.EXCEL_PATH.with_suffix(".xlsx.tmp")
    try:
        workbook.save(tmp_path)
        tmp_path.replace(config.EXCEL_PATH)
    except PermissionError:
        os.replace(dest, src)  # roll back the file move, sheet was never written
        return False, "locked"
    return True, None


def _bar_row(label: str, count: int, max_count: int) -> str:
    pct = round(100 * count / max_count) if max_count else 0
    return f"""
    <div class="hbar-row">
      <div class="hbar-label">{html.escape(label)}</div>
      <div class="hbar-track"><div class="hbar-fill" style="width:{pct}%"></div></div>
      <div class="hbar-value">{count}</div>
    </div>"""


def _hour_bar(hour: int, entries: list[dict], max_count: int) -> str:
    count = len(entries)
    height_pct = round(100 * count / max_count) if max_count else 0
    return f"""
    <div class="vbar-col">
      <div class="vbar-track">
        <div class="vbar-fill" style="height:{height_pct}%"></div>
      </div>
      <div class="vbar-count">{count if count else ""}</div>
      <div class="vbar-label">{hour:02d}</div>
    </div>"""


def _table_rows(rows: list[dict], show_archive_button: bool, date_range: str = "today") -> str:
    out = []
    for r in reversed(rows):  # most recent first
        plate = r["plate"] or "unreadable"
        make_model = r["make_model"] or "unknown"
        src = f"/snapshots/{r['snapshot']}" if r["snapshot"] else ""
        thumb = f'<a href="{html.escape(src)}" target="_blank"><img class="thumb" src="{html.escape(src)}" loading="lazy" /></a>' if src else ""
        action = ""
        if show_archive_button and r["snapshot"]:
            action = f"""<form method="post" action="/archive" class="inline-form">
              <input type="hidden" name="snapshot" value="{html.escape(r['snapshot'])}" />
              <input type="hidden" name="date_range" value="{html.escape(date_range)}" />
              <button type="submit" class="archive-btn">Archive</button>
            </form>"""
        out.append(f"""
        <tr>
          <td>{thumb}</td>
          <td class="mono">{html.escape(r["date"])}</td>
          <td class="mono">{html.escape(r["time"])}</td>
          <td>{html.escape(r["vehicle_type"] or "")}</td>
          <td>{html.escape(make_model)}</td>
          <td class="mono">{html.escape(plate)}</td>
          <td>{action}</td>
        </tr>""")
    return "".join(out)


_ERROR_MESSAGES = {
    "locked": "Could not save — detections.xlsx is open in another program (e.g. Excel). Close it and try again.",
    "not_found": "Could not find that detection — the page may be out of date, try refreshing.",
}


def render_dashboard(tab: str = "cars", date_range: str = "today", error: str | None = None) -> str:
    tab = tab if tab in ("cars", "archive") else "cars"
    date_range = date_range if date_range in _RANGE_LABELS else "today"
    error_banner = (
        f'<div class="error-banner">{html.escape(_ERROR_MESSAGES.get(error, error))}</div>' if error else ""
    )

    start, end = _date_bounds(date_range)
    rows = [r for r in _read_rows() if _in_range(r, start, end)]

    cars_rows = [r for r in rows if not _is_archived(r)]
    archive_rows = [r for r in rows if _is_archived(r)]
    active_rows = cars_rows if tab == "cars" else archive_rows

    manufacturer_counts = Counter(m for r in rows if (m := _manufacturer(r["make_model"])))
    top_manufacturers = manufacturer_counts.most_common(5)
    top_max = top_manufacturers[0][1] if top_manufacturers else 1

    hours: dict[int, list[dict]] = defaultdict(list)
    for r in rows:
        hour = int(r["time"].split(":")[0])
        hours[hour].append(r)
    hour_max = max((len(v) for v in hours.values()), default=0) or 1

    range_buttons = "".join(
        f'<a class="range-btn {"active" if key == date_range else ""}" href="/dashboard?tab={tab}&date_range={key}">{label}</a>'
        for key, label in _RANGE_LABELS.items()
    )
    qs = f"date_range={date_range}"

    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8" />
<title>Watch Street Cars</title>
<style>
  .viz-root {{
    color-scheme: light;
    --page: #ffffff;
    --surface: #ffffff;
    --text-primary: #000000;
    --text-secondary: #000000;
    --text-muted: #000000;
    --gridline: #e1e0d9;
    --baseline: #c3c2b7;
    --border: rgba(11,11,11,0.10);
    --series-1: #2a78d6;
  }}
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; font-family: system-ui, -apple-system, "Segoe UI", sans-serif; background: var(--page); color: var(--text-primary); }}
  .page-header {{ padding: 16px 16px 0; }}
  .page-title {{ font-size: 22px; font-weight: 700; margin: 0; }}
  .error-banner {{ margin: 12px 16px 0; padding: 10px 14px; border-radius: 6px; background: #fdeaea; border: 1px solid #e34948; color: #8a1c1c; font-size: 13px; }}
  .layout {{ display: flex; gap: 16px; padding: 16px; height: calc(100vh - 46px); }}
  .panel {{ background: var(--surface); border: 1px solid var(--border); border-radius: 8px; padding: 16px; overflow: auto; }}
  .left {{ flex: 1.4; min-width: 0; display: flex; flex-direction: column; }}
  .right {{ flex: 1; min-width: 320px; display: flex; flex-direction: column; gap: 16px; }}
  h1 {{ font-size: 18px; margin: 0 0 12px; }}
  h2 {{ font-size: 14px; margin: 0 0 12px; color: var(--text-secondary); font-weight: 600; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
  thead th {{ position: sticky; top: 0; background: var(--surface); text-align: left; padding: 8px 6px; color: var(--text-muted); font-weight: 600; border-bottom: 1px solid var(--gridline); }}
  tbody td {{ padding: 6px; border-bottom: 1px solid var(--gridline); vertical-align: middle; }}
  .mono {{ font-variant-numeric: tabular-nums; color: var(--text-secondary); }}
  .thumb {{ width: 48px; height: 36px; object-fit: cover; border-radius: 4px; display: block; }}
  .tabs {{ display: flex; gap: 4px; margin-bottom: 12px; border-bottom: 1px solid var(--gridline); }}
  .tab-link {{ padding: 8px 14px; text-decoration: none; color: var(--text-secondary); font-size: 13px; font-weight: 600; border-bottom: 2px solid transparent; margin-bottom: -1px; }}
  .tab-link.active {{ color: var(--series-1); border-bottom-color: var(--series-1); }}
  .inline-form {{ margin: 0; }}
  .archive-btn {{ font-family: inherit; font-size: 12px; padding: 4px 10px; border-radius: 6px; border: 1px solid var(--border); background: var(--surface); color: var(--text-primary); cursor: pointer; }}
  .archive-btn:hover {{ background: var(--gridline); }}
  .range-buttons {{ display: flex; flex-wrap: wrap; gap: 8px; }}
  .range-btn {{ font-family: inherit; font-size: 13px; font-weight: 600; padding: 8px 14px; border-radius: 6px; border: 1px solid var(--border); background: var(--surface); color: var(--text-primary); cursor: pointer; text-decoration: none; }}
  .range-btn:hover {{ background: var(--gridline); }}
  .range-btn.active {{ background: var(--series-1); border-color: var(--series-1); color: #ffffff; }}
  .hbar-row {{ display: flex; align-items: center; gap: 8px; margin-bottom: 10px; }}
  .hbar-label {{ width: 110px; font-size: 13px; color: var(--text-secondary); flex-shrink: 0; }}
  .hbar-track {{ flex: 1; background: var(--gridline); border-radius: 4px; height: 20px; }}
  .hbar-fill {{ height: 100%; background: var(--series-1); border-radius: 4px; }}
  .hbar-value {{ width: 24px; text-align: right; font-size: 13px; color: var(--text-secondary); font-variant-numeric: tabular-nums; }}
  .vbar-chart {{ display: flex; align-items: flex-end; gap: 2px; height: 160px; border-bottom: 1px solid var(--baseline); }}
  .vbar-col {{ flex: 1; display: flex; flex-direction: column; align-items: center; justify-content: flex-end; height: 100%; position: relative; }}
  .vbar-track {{ width: 100%; max-width: 18px; height: 130px; display: flex; align-items: flex-end; flex-shrink: 0; }}
  .vbar-fill {{ width: 100%; background: var(--series-1); border-radius: 3px 3px 0 0; min-height: 0; }}
  .vbar-count {{ font-size: 10px; color: var(--text-muted); height: 12px; font-variant-numeric: tabular-nums; }}
  .vbar-label {{ font-size: 10px; color: var(--text-muted); margin-top: 4px; }}
  .empty {{ color: var(--text-muted); font-size: 13px; }}
</style>
</head>
<body>
<div class="viz-root">
  <div class="page-header">
    <h1 class="page-title">Watch Street Cars</h1>
  </div>
  {error_banner}
  <div class="layout">
    <div class="panel left">
      <div class="tabs">
        <a class="tab-link {'active' if tab == 'cars' else ''}" href="/dashboard?tab=cars&{qs}">Cars ({len(cars_rows)})</a>
        <a class="tab-link {'active' if tab == 'archive' else ''}" href="/dashboard?tab=archive&{qs}">Archive ({len(archive_rows)})</a>
      </div>
      <table>
        <thead><tr><th></th><th>Date</th><th>Time</th><th>Type</th><th>Make / Model</th><th>Plate</th><th></th></tr></thead>
        <tbody>{_table_rows(active_rows, show_archive_button=(tab == "cars"), date_range=date_range) if active_rows else f'<tr><td colspan="7" class="empty">No {"cars" if tab == "cars" else "archived"} detections in this range.</td></tr>'}</tbody>
      </table>
    </div>
    <div class="right">
      <div class="panel">
        <h2>Dates</h2>
        <div class="range-buttons">{range_buttons}</div>
      </div>
      <div class="panel">
        <h2>Top 5 manufacturers ({_RANGE_LABELS[date_range]})</h2>
        {"".join(_bar_row(name, count, top_max) for name, count in top_manufacturers) if top_manufacturers else '<div class="empty">No identified manufacturers in this range.</div>'}
      </div>
      <div class="panel">
        <h2>Hourly timeline ({_RANGE_LABELS[date_range]})</h2>
        {'<div class="vbar-chart">' + "".join(_hour_bar(h, hours.get(h, []), hour_max) for h in range(24)) + '</div>' if rows else '<div class="empty">No detections in this range.</div>'}
      </div>
    </div>
  </div>
</div>
</body>
</html>"""
