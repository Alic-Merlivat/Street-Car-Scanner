"""Re-runs the (now Claude-vision-based) make/model + plate recognition
over every image already sitting in data/snapshots (not review/), and
brings detections.xlsx up to date:
  - existing rows for those files get their Make/Model + Plate columns
    refreshed with the better reading
  - files with no row at all (lost during an earlier crash before the
    atomic-save fix landed) get a new row added, thumbnail included

Doesn't touch data/snapshots/review - those are flagged as not showing
the whole car, so a better plate/model reading wouldn't be reliable anyway.
"""
import datetime as dt
import io
from pathlib import Path

import cv2
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

load_dotenv(Path(__file__).resolve().parent / ".env")

import config
from vehicle_detector import VehicleDetector
from vision_classifier import VisionClassifier

_SNAPSHOT_COLUMN = 8
_SNAPSHOT_COLUMN_LETTER = "H"
_THUMBNAIL_WIDTH_PX = 120
_ROW_HEIGHT_POINTS = 70


def make_thumbnail(snapshot_bgr):
    height, width = snapshot_bgr.shape[:2]
    if width == 0 or height == 0:
        return None
    scale = _THUMBNAIL_WIDTH_PX / width
    thumb = cv2.resize(snapshot_bgr, (_THUMBNAIL_WIDTH_PX, max(1, round(height * scale))))
    ok, encoded = cv2.imencode(".png", thumb)
    return io.BytesIO(encoded.tobytes()) if ok else None


def timestamp_from_filename(filename: str) -> dt.datetime:
    stem = filename.rsplit(".", 1)[0]  # 20260917_204306_517775
    date_part, time_part, _ = stem.split("_")
    return dt.datetime.strptime(date_part + time_part, "%Y%m%d%H%M%S")


def main() -> None:
    files = sorted(p for p in config.SNAPSHOT_DIR.glob("*.jpg"))
    print(f"reprocessing {len(files)} image(s) in {config.SNAPSHOT_DIR}")

    detector = VehicleDetector()
    classifier = VisionClassifier()

    workbook = load_workbook(config.EXCEL_PATH)
    sheet = workbook["Detections"]
    row_by_filename = {row[7].value: row[7].row for row in sheet.iter_rows(min_row=2) if row[7].value}

    for path in files:
        image = cv2.imread(str(path))
        if image is None:
            print(f"skip (unreadable): {path.name}")
            continue

        vehicles = detector.detect(image)
        vehicle_label = max(vehicles, key=lambda v: v.confidence).label if vehicles else "car"

        result = classifier.classify(image)
        print(f"{path.name}: {vehicle_label} | {result.make_model} | {result.plate}")

        make_model_confidence = 1.0 if result.make_model else 0.0
        plate_confidence = 1.0 if result.plate else 0.0

        if path.name in row_by_filename:
            row = row_by_filename[path.name]
            sheet.cell(row=row, column=3, value=vehicle_label)
            sheet.cell(row=row, column=4, value=result.make_model or "unknown")
            sheet.cell(row=row, column=5, value=round(make_model_confidence, 2))
            sheet.cell(row=row, column=6, value=result.plate or "unreadable")
            sheet.cell(row=row, column=7, value=round(plate_confidence, 2))
        else:
            timestamp = timestamp_from_filename(path.name)
            sheet.append([
                timestamp.strftime("%Y-%m-%d"),
                timestamp.strftime("%H:%M:%S"),
                vehicle_label,
                result.make_model or "unknown",
                round(make_model_confidence, 2),
                result.plate or "unreadable",
                round(plate_confidence, 2),
                path.name,
            ])
            row = sheet.max_row
            sheet.row_dimensions[row].height = _ROW_HEIGHT_POINTS
            sheet.column_dimensions[_SNAPSHOT_COLUMN_LETTER].width = 24
            cell = sheet.cell(row=row, column=_SNAPSHOT_COLUMN)
            cell.hyperlink = path.as_uri()
            cell.style = "Hyperlink"
            thumbnail = make_thumbnail(image)
            if thumbnail is not None:
                img = XLImage(thumbnail)
                img.anchor = f"{_SNAPSHOT_COLUMN_LETTER}{row}"
                sheet.add_image(img)

    tmp_path = config.EXCEL_PATH.with_suffix(".xlsx.tmp")
    workbook.save(tmp_path)
    tmp_path.replace(config.EXCEL_PATH)
    print("done")


if __name__ == "__main__":
    main()
