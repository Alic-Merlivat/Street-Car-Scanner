"""Appends detection events to a local .xlsx file, one row per car."""
import datetime as dt
import io
import os
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage

import config

_HEADER = [
    "Date",
    "Time",
    "Vehicle Type",
    "Make / Model",
    "Make/Model Confidence",
    "Plate Number",
    "Plate Confidence",
    "Snapshot File",
]

_SNAPSHOT_COLUMN = 8  # H
_SNAPSHOT_COLUMN_LETTER = "H"
_THUMBNAIL_WIDTH_PX = 120
_ROW_HEIGHT_POINTS = 70  # ~93px, enough for a ~120px-wide thumbnail at typical aspect ratios


class ExcelExporter:
    def __init__(self) -> None:
        config.DATA_DIR.mkdir(parents=True, exist_ok=True)
        config.SNAPSHOT_DIR.mkdir(parents=True, exist_ok=True)
        config.REVIEW_DIR.mkdir(parents=True, exist_ok=True)
        if not config.EXCEL_PATH.exists():
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Detections"
            sheet.append(_HEADER)
            self._save(workbook)

    @staticmethod
    def _save(workbook: Workbook) -> None:
        """Writes to a temp file then renames it into place, so a crash or
        forced kill mid-save can never leave detections.xlsx corrupted -
        the rename is atomic, the old file stays intact until it succeeds."""
        tmp_path = config.EXCEL_PATH.with_suffix(".xlsx.tmp")
        workbook.save(tmp_path)
        os.replace(tmp_path, config.EXCEL_PATH)

    @staticmethod
    def _make_thumbnail(snapshot_bgr: np.ndarray) -> io.BytesIO | None:
        height, width = snapshot_bgr.shape[:2]
        if width == 0 or height == 0:
            return None
        scale = _THUMBNAIL_WIDTH_PX / width
        thumb = cv2.resize(snapshot_bgr, (_THUMBNAIL_WIDTH_PX, max(1, round(height * scale))))
        encode_ok, encoded = cv2.imencode(".png", thumb)
        if not encode_ok:
            return None
        return io.BytesIO(encoded.tobytes())

    def append_event(
        self,
        vehicle_label: str,
        make_model: str | None,
        make_model_confidence: float,
        plate_text: str | None,
        plate_confidence: float,
        snapshot_bgr: np.ndarray,
        needs_review: bool = False,
    ) -> Path:
        now = dt.datetime.now()
        snapshot_name = f"{now.strftime('%Y%m%d_%H%M%S_%f')}.jpg"
        snapshot_dir = config.REVIEW_DIR if needs_review else config.SNAPSHOT_DIR
        snapshot_path = snapshot_dir / snapshot_name
        cv2.imwrite(str(snapshot_path), snapshot_bgr)
        display_name = str(snapshot_path.relative_to(config.SNAPSHOT_DIR)).replace("\\", "/")

        workbook = load_workbook(config.EXCEL_PATH)
        sheet = workbook["Detections"]
        sheet.append(
            [
                now.strftime("%Y-%m-%d"),
                now.strftime("%H:%M:%S"),
                vehicle_label,
                make_model or "unknown",
                round(make_model_confidence, 2),
                plate_text or "unreadable",
                round(plate_confidence, 2),
                display_name,
            ]
        )
        row = sheet.max_row
        sheet.row_dimensions[row].height = _ROW_HEIGHT_POINTS
        sheet.column_dimensions[_SNAPSHOT_COLUMN_LETTER].width = 24

        cell = sheet.cell(row=row, column=_SNAPSHOT_COLUMN)
        cell.hyperlink = snapshot_path.as_uri()
        cell.style = "Hyperlink"

        thumbnail = self._make_thumbnail(snapshot_bgr)
        if thumbnail is not None:
            image = XLImage(thumbnail)
            image.anchor = f"{_SNAPSHOT_COLUMN_LETTER}{row}"
            sheet.add_image(image)

        self._save(workbook)
        return snapshot_path
