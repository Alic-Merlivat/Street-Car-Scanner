"""One-off cleanup: moves already-saved snapshots that don't show the whole
vehicle into data/snapshots/review, based on a manual visual review (there's
no bounding-box/frame data saved for these older crops to check
automatically - that check now happens live in app.py for anything captured
from now on). Also repoints any existing Excel rows that reference a moved
file so their links don't break.
"""
from pathlib import Path

from openpyxl import load_workbook

import config

REVIEW_FILENAMES = {
    "20260917_204209_686542.jpg", "20260917_204233_353203.jpg", "20260917_204233_966090.jpg",
    "20260917_204238_349107.jpg", "20260917_204238_768001.jpg", "20260917_204239_486792.jpg",
    "20260917_204239_895507.jpg", "20260917_204256_073218.jpg", "20260917_204256_483555.jpg",
    "20260917_204259_630944.jpg", "20260917_204300_449747.jpg", "20260917_204300_834767.jpg",
    "20260917_204302_292789.jpg", "20260917_204310_557154.jpg", "20260917_204311_656345.jpg",
    "20260917_204314_735327.jpg", "20260917_204315_876775.jpg", "20260917_211118_757531.jpg",
    "20260917_211126_517881.jpg", "20260917_211132_401170.jpg", "20260917_211140_593093.jpg",
    "20260917_211141_397001.jpg", "20260917_211141_976401.jpg", "20260917_211145_934933.jpg",
    "20260917_211148_160767.jpg", "20260917_211152_241312.jpg", "20260917_211152_704836.jpg",
    "20260917_211156_708356.jpg", "20260917_211203_161978.jpg", "20260917_211203_612549.jpg",
}


def main() -> None:
    config.REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    moved = []
    for filename in sorted(REVIEW_FILENAMES):
        src = config.SNAPSHOT_DIR / filename
        if not src.exists():
            print(f"skip (not found): {filename}")
            continue
        dest = config.REVIEW_DIR / filename
        src.rename(dest)
        moved.append(filename)
    print(f"moved {len(moved)} file(s) into {config.REVIEW_DIR}")

    if not config.EXCEL_PATH.exists():
        return
    workbook = load_workbook(config.EXCEL_PATH)
    sheet = workbook["Detections"]
    updated_rows = 0
    for row in sheet.iter_rows(min_row=2):
        cell = row[7]  # Snapshot File column
        if cell.value in REVIEW_FILENAMES:
            cell.value = f"review/{cell.value}"
            new_path = config.REVIEW_DIR / cell.value.split("/", 1)[1]
            if new_path.exists():
                cell.hyperlink = new_path.as_uri()
            updated_rows += 1
    if updated_rows:
        tmp_path = config.EXCEL_PATH.with_suffix(".xlsx.tmp")
        workbook.save(tmp_path)
        tmp_path.replace(config.EXCEL_PATH)
    print(f"updated {updated_rows} existing Excel row(s) to point at the new location")


if __name__ == "__main__":
    main()
